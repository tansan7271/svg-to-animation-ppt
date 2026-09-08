#!/usr/bin/env python3
"""svg2inkppt: SVG 파일을 PowerPoint 잉크(Ink) 개체로 변환하고 '재생' 애니메이션을 붙인다.

사용법:
    python svg2inkppt.py logo.svg                      # logo.pptx 생성
    python svg2inkppt.py logo.svg -o out.pptx --width-cm 12 --duration 3

동작 원리:
    SVG의 path/rect/circle/polyline 등을 점열로 샘플링해 InkML <trace>로 만들고,
    슬라이드에 <p:contentPart>로 심은 뒤, drawProgress 0→1 애니메이션(재생)을 건다.
"""
from __future__ import annotations

import argparse
import copy
import datetime as _dt
import math
import sys
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree

_XP = etree.XMLParser(remove_blank_text=True)
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu
from svgelements import (
    SVG, Path as SvgPath, Shape, Group, Move, Close, Color, Matrix,
)

EMU_PER_CM = 360_000
INK_PER_CM = 1000  # InkML resolution 1000/cm  (1 unit = 0.001 cm)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
RT_CUSTOM_XML = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"


# ----------------------------------------------------------------------------
# 1. SVG -> 획(stroke) 리스트
# ----------------------------------------------------------------------------
@dataclass
class Stroke:
    points: list[tuple[float, float]]  # SVG user units (변환 적용 후)
    color: str                          # "#RRGGBB"
    width: float                        # SVG user units
    order: int = 0                      # 문서 순서 (path 단위 번호)
    name: str = ""                      # SVG id 또는 태그명


def _color_hex(c) -> str | None:
    if c is None:
        return None
    try:
        col = Color(c)
    except Exception:
        return None
    if col.value is None or col.alpha == 0:
        return None
    return "#%02X%02X%02X" % (col.red, col.green, col.blue)


def _sample_subpath(sub: SvgPath, step: float) -> list[tuple[float, float]]:
    """arc-length 기준으로 step 간격 샘플링. step은 SVG user unit."""
    pts: list[tuple[float, float]] = []
    for seg in sub:
        if isinstance(seg, Move):
            pts.append((seg.end.x, seg.end.y))
            continue
        if isinstance(seg, Close):
            if seg.end is not None and seg.start is not None:
                pts.append((seg.end.x, seg.end.y))
            continue
        try:
            length = seg.length(error=1e-4)
        except Exception:
            length = 0
        n = max(1, int(math.ceil(length / step)))
        for i in range(1, n + 1):
            p = seg.point(i / n)
            pts.append((p.x, p.y))
    # 중복 제거
    out: list[tuple[float, float]] = []
    for p in pts:
        if not out or abs(out[-1][0] - p[0]) > 1e-6 or abs(out[-1][1] - p[1]) > 1e-6:
            out.append(p)
    return out


def load_strokes(svg_file: Path, step_units: float | None, default_width_units: float | None):
    svg = SVG.parse(str(svg_file), reify=True)
    strokes: list[Stroke] = []
    order = 0

    # SVG 전체 크기(user units)를 이용해 기본 샘플 간격 결정
    vb_w = float(svg.width) if svg.width else 100.0
    vb_h = float(svg.height) if svg.height else 100.0
    if step_units is None:
        step_units = max(vb_w, vb_h) / 600.0
    if default_width_units is None:
        default_width_units = max(vb_w, vb_h) / 200.0

    for el in svg.elements():
        if not isinstance(el, Shape):
            continue
        # 화면에 안 보이는 것 제외
        if getattr(el, "values", {}).get("visibility") == "hidden":
            continue
        if getattr(el, "values", {}).get("display") == "none":
            continue
        path = SvgPath(el) if not isinstance(el, SvgPath) else el
        try:
            path.reify()
        except Exception:
            pass
        el_id = getattr(el, "id", None) or getattr(el, "values", {}).get("id")
        name = str(el_id) if el_id else type(el).__name__.lower()
        stroke_hex = _color_hex(el.stroke)
        fill_hex = _color_hex(el.fill)
        color = stroke_hex or fill_hex or "#000000"
        if stroke_hex and el.stroke_width:
            width = float(el.stroke_width)
        else:
            width = default_width_units
        added = False
        for sub in path.as_subpaths():
            sub_path = SvgPath(sub)
            pts = _sample_subpath(sub_path, step_units)
            if len(pts) < 2:
                continue
            strokes.append(Stroke(points=pts, color=color, width=width, order=order, name=name))
            added = True
        if added:
            order += 1
    if not strokes:
        sys.exit("SVG에서 그릴 수 있는 path/도형을 찾지 못했어요.")
    return strokes


def apply_order(strokes: list[Stroke], path_order: list[int] | None,
                reversed_paths: set[int] | None = None) -> list[Stroke]:
    """path_order: 원래 order 번호들의 새 순서. reversed_paths: 방향을 뒤집을 원래 order 번호들.
    반환: 새 순서로 정렬되고 order 가 0.. 으로 다시 매겨진 획 목록 (원본은 건드리지 않음)."""
    reversed_paths = reversed_paths or set()
    orig_orders = sorted({s.order for s in strokes})
    seq = list(path_order) if path_order else orig_orders
    seq += [o for o in orig_orders if o not in seq]          # 빠진 건 뒤에
    rank = {o: i for i, o in enumerate(seq)}
    out: list[Stroke] = []
    for o in seq:
        grp = [s for s in strokes if s.order == o]
        if o in reversed_paths:
            grp = [Stroke(points=list(reversed(s.points)), color=s.color, width=s.width,
                          order=rank[o], name=s.name) for s in reversed(grp)]
        else:
            grp = [Stroke(points=list(s.points), color=s.color, width=s.width,
                          order=rank[o], name=s.name) for s in grp]
        out += grp
    return out


# ----------------------------------------------------------------------------
# 2. 획 -> InkML
# ----------------------------------------------------------------------------
@dataclass
class InkResult:
    xml: bytes
    width_cm: float
    height_cm: float
    strokes_cm: list[Stroke]   # 좌표가 cm 로 변환된 획(미리보기 PNG 용)
    brush_widths_cm: dict[str, float] = field(default_factory=dict)


def build_inkml(strokes: list[Stroke], target_width_cm: float, target_height_cm: float | None,
                stroke_width_cm: float | None, color_override: str | None,
                timing: str = "channel", speed_cm_s: float = 10.0, gap_ms: float = 60.0,
                scale: float | None = None) -> InkResult:
    """timing: none   = 시각 정보 없음
               offset = trace 마다 timeOffset 만
               channel= timeOffset + T 채널(점마다 ms)"""
    xs = [x for s in strokes for x, _ in s.points]
    ys = [y for s in strokes for _, y in s.points]
    minx, maxx, miny, maxy = min(xs), max(xs), min(ys), max(ys)
    w_units = max(maxx - minx, 1e-9)
    h_units = max(maxy - miny, 1e-9)

    if scale is None:
        scale = target_width_cm / w_units
        if target_height_cm and h_units * scale > target_height_cm:
            scale = target_height_cm / h_units

    width_cm = w_units * scale
    height_cm = h_units * scale

    # 붓(brush)은 (색, 굵기) 조합마다 하나
    brushes: dict[tuple[str, float], str] = {}
    strokes_cm: list[Stroke] = []
    traces = []
    for s in strokes:
        color = color_override or s.color
        w_cm = stroke_width_cm if stroke_width_cm else max(0.01, round(s.width * scale, 3))
        key = (color, w_cm)
        if key not in brushes:
            brushes[key] = f"br{len(brushes)}"
        pts_cm = [((x - minx) * scale, (y - miny) * scale) for x, y in s.points]
        strokes_cm.append(Stroke(points=pts_cm, color=color, width=w_cm, order=s.order))
        traces.append((brushes[key], pts_cm))

    # 길이 비례 시각: 획 시작 timeOffset(ms) + 점마다 경과 ms
    t_cursor = 0.0
    encoded = []
    for bid, pts_cm in traces:
        seg_len = [math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(pts_cm, pts_cm[1:])]
        cum = [0.0]
        for L in seg_len:
            cum.append(cum[-1] + L)
        times_ms = [c / speed_cm_s * 1000.0 for c in cum]
        encoded.append((bid, t_cursor, times_ms, pts_cm))
        t_cursor += times_ms[-1] + gap_ms

    ts = _dt.datetime.now().strftime("%Y-%m-%dT%H:%M:%S.000")
    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<inkml:ink xmlns:inkml="http://www.w3.org/2003/InkML">',
        '<inkml:definitions>',
        '<inkml:context xml:id="ctx0"><inkml:inkSource xml:id="inkSrc0"><inkml:traceFormat>'
        '<inkml:channel name="X" type="integer" min="-2.14748E9" max="2.14748E9" units="cm"/>'
        '<inkml:channel name="Y" type="integer" min="-2.14748E9" max="2.14748E9" units="cm"/>'
        '<inkml:channel name="F" type="integer" max="32767" units="dev"/>'
        + ('<inkml:channel name="T" type="integer" max="2.14748E9" units="dev"/>' if timing == "channel" else '')
        + '</inkml:traceFormat><inkml:channelProperties>'
        '<inkml:channelProperty channel="X" name="resolution" value="1000" units="1/cm"/>'
        '<inkml:channelProperty channel="Y" name="resolution" value="1000" units="1/cm"/>'
        '<inkml:channelProperty channel="F" name="resolution" value="0" units="1/dev"/>'
        + ('<inkml:channelProperty channel="T" name="resolution" value="1" units="1/dev"/>' if timing == "channel" else '')
        + '</inkml:channelProperties></inkml:inkSource>'
        f'<inkml:timestamp xml:id="ts0" timeString="{ts}"/></inkml:context>',
    ]
    for (color, w_cm), bid in brushes.items():
        parts.append(
            f'<inkml:brush xml:id="{bid}">'
            f'<inkml:brushProperty name="width" value="{w_cm}" units="cm"/>'
            f'<inkml:brushProperty name="height" value="{w_cm}" units="cm"/>'
            f'<inkml:brushProperty name="color" value="{color}"/>'
            f'<inkml:brushProperty name="ignorePressure" value="1"/>'
            '</inkml:brush>'
        )
    parts.append('</inkml:definitions>')
    for bid, t0, times_ms, pts_cm in encoded:
        attrs = f'contextRef="#ctx0" brushRef="#{bid}"'
        if timing in ("offset", "channel"):
            attrs += f' timeOffset="{int(round(t0))}"'
        body = _encode_trace(pts_cm, times_ms if timing == "channel" else None)
        parts.append(f'<inkml:trace {attrs}>{body}</inkml:trace>')
    parts.append('</inkml:ink>')
    xml = "".join(parts).encode("utf-8")
    return InkResult(xml=xml, width_cm=width_cm, height_cm=height_cm, strokes_cm=strokes_cm,
                     brush_widths_cm={bid: w for (c, w), bid in brushes.items()})


def _encode_trace(pts_cm: list[tuple[float, float]], times_ms: list[float] | None = None,
                  pressure: int = 16384) -> str:
    """PowerPoint 식 InkML 인코딩: 첫 샘플 절대값, 이후 ' 접두어 1차 차분."""
    ints = [(int(round(x * INK_PER_CM)), int(round(y * INK_PER_CM))) for x, y in pts_cm]
    ts = [int(round(t)) for t in times_ms] if times_ms else None
    first = f"{ints[0][0]} {ints[0][1]} {pressure}" + (f" {ts[0]}" if ts else "")
    out = [first]
    px, py = ints[0]
    pt = ts[0] if ts else 0
    for i, (x, y) in enumerate(ints[1:], start=1):
        item = f"'{x - px}'{y - py}'0"
        if ts:
            item += f"'{ts[i] - pt}"
            pt = ts[i]
        out.append(item)
        px, py = x, y
    return ",".join(out)


# ----------------------------------------------------------------------------
# 3. 미리보기 PNG (구버전/타 앱용 fallback 그림)
# ----------------------------------------------------------------------------
def render_fallback_png(ink: InkResult, out_file: Path, px_per_cm: float = 60.0):
    margin_cm = max(ink.brush_widths_cm.values()) if ink.brush_widths_cm else 0.05
    ss = 3  # supersampling
    W = int(math.ceil((ink.width_cm + 2 * margin_cm) * px_per_cm)) * ss
    H = int(math.ceil((ink.height_cm + 2 * margin_cm) * px_per_cm)) * ss
    img = Image.new("RGBA", (max(W, 1), max(H, 1)), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    for s in ink.strokes_cm:
        pts = [((x + margin_cm) * px_per_cm * ss, (y + margin_cm) * px_per_cm * ss) for x, y in s.points]
        w = max(1, int(round(s.width * px_per_cm * ss)))
        r = w / 2
        d.line(pts, fill=s.color, width=w)
        for (x, y) in pts:
            d.ellipse([x - r, y - r, x + r, y + r], fill=s.color)
    img = img.resize((W // ss, H // ss), Image.LANCZOS)
    img.save(out_file)
    return margin_cm


# ----------------------------------------------------------------------------
# 4. pptx 조립
# ----------------------------------------------------------------------------
CONTENT_PART_XML = """
<mc:AlternateContent xmlns:mc="{mc}" xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}">
  <mc:Choice xmlns:p14="{p14}" Requires="p14">
    <p:contentPart p14:bwMode="auto" r:id="{ink_rid}">
      <p14:nvContentPartPr>
        <p14:cNvPr id="{sid}" name="{name}"/>
        <p14:cNvContentPartPr/>
        <p14:nvPr/>
      </p14:nvContentPartPr>
      <p14:xfrm>
        <a:off x="{off_x}" y="{off_y}"/>
        <a:ext cx="{ext_cx}" cy="{ext_cy}"/>
      </p14:xfrm>
    </p:contentPart>
  </mc:Choice>
  <mc:Fallback>
    <p:pic>
      <p:nvPicPr>
        <p:cNvPr id="{sid}" name="{name}"/>
        <p:cNvPicPr/>
        <p:nvPr/>
      </p:nvPicPr>
      <p:blipFill>
        <a:blip r:embed="{img_rid}"/>
        <a:stretch><a:fillRect/></a:stretch>
      </p:blipFill>
      <p:spPr>
        <a:xfrm>
          <a:off x="{pic_off_x}" y="{pic_off_y}"/>
          <a:ext cx="{pic_ext_cx}" cy="{pic_ext_cy}"/>
        </a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      </p:spPr>
    </p:pic>
  </mc:Fallback>
</mc:AlternateContent>
"""

# 잉크 '재생' 효과 한 개 (presetID 63, drawProgress 0 -> 1)
INK_EFFECT_XML = """
<p:par xmlns:p="{p}">
  <p:cTn id="{id0}" presetID="63" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node_type}">
    <p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>
    <p:childTnLst>
      <p:set>
        <p:cBhvr>
          <p:cTn id="{id1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
          <p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:to><p:strVal val="visible"/></p:to>
      </p:set>
      <p:anim calcmode="lin" valueType="num">
        <p:cBhvr>
          <p:cTn id="{id2}" dur="{dur}" fill="hold"{accel_attrs}/>
          <p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>drawProgress</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:tavLst>{tavs}</p:tavLst>
      </p:anim>
    </p:childTnLst>
  </p:cTn>
</p:par>
"""

TIMING_XML = """
<p:timing xmlns:p="{p}">
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                <p:par>
                  <p:cTn id="3" fill="hold">
                    <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                    <p:childTnLst>
                      <p:par>
                        <p:cTn id="4" fill="hold">
                          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                          <p:childTnLst/>
                        </p:cTn>
                      </p:par>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
            <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>
"""


def ease_curve(ease_in: float, ease_out: float, n: int = 60) -> list[tuple[float, float]]:
    """PowerPoint accel/decel 과 같은 모델: 앞 ease_in 구간 등가속, 뒤 ease_out 구간 등감속.
    ease_in + ease_out <= 1. 반환값: (시간 0..1, 진행률 0..1) 키프레임 목록."""
    a = max(0.0, min(1.0, ease_in))
    d = max(0.0, min(1.0 - a, ease_out))
    # 최고 속도 v: 면적 = v*(a/2) + v*(1-a-d) + v*(d/2) = 1
    v = 1.0 / (1.0 - a / 2 - d / 2) if (1.0 - a / 2 - d / 2) > 0 else 1.0

    def prog(t: float) -> float:
        if a > 0 and t < a:
            return v * t * t / (2 * a)
        p_a = v * a / 2
        if t < 1 - d:
            return p_a + v * (t - a)
        p_b = p_a + v * (1 - d - a)
        u = t - (1 - d)
        return p_b + v * u - v * u * u / (2 * d) if d > 0 else p_b

    pts = [(i / n, prog(i / n)) for i in range(n + 1)]
    pts[-1] = (1.0, 1.0)
    return pts


def _tav_xml(curve: list[tuple[float, float]]) -> str:
    return "".join(
        f'<p:tav tm="{int(round(t * 100000))}"><p:val><p:fltVal val="{v:.5f}"/></p:val></p:tav>'
        for t, v in curve
    )


def _next_ink_partname(package) -> PackURI:
    existing = {p.partname for p in package.iter_parts()}
    n = 1
    while PackURI(f"/ppt/ink/ink{n}.xml") in existing:
        n += 1
    return PackURI(f"/ppt/ink/ink{n}.xml")


def _next_shape_id(slide) -> int:
    ids = [int(v) for v in slide._element.xpath("//@id") if str(v).isdigit()]
    return max(ids + [1]) + 1


def add_ink_to_slide(slide, ink: InkResult, png_file: Path, margin_cm: float,
                     off_x_emu: int, off_y_emu: int, name: str = "Ink 1") -> int:
    package = slide.part.package
    ink_part = Part(_next_ink_partname(package), "application/inkml+xml", package, ink.xml)
    ink_rid = slide.part.relate_to(ink_part, RT_CUSTOM_XML)
    image_part, img_rid = slide.part.get_or_add_image_part(str(png_file))

    sid = _next_shape_id(slide)
    m = int(round(margin_cm * EMU_PER_CM))
    xml = CONTENT_PART_XML.format(
        **NS, ink_rid=ink_rid, img_rid=img_rid, sid=sid, name=name,
        off_x=off_x_emu, off_y=off_y_emu,
        ext_cx=int(round(ink.width_cm * EMU_PER_CM)), ext_cy=int(round(ink.height_cm * EMU_PER_CM)),
        pic_off_x=off_x_emu - m, pic_off_y=off_y_emu - m,
        pic_ext_cx=int(round((ink.width_cm + 2 * margin_cm) * EMU_PER_CM)),
        pic_ext_cy=int(round((ink.height_cm + 2 * margin_cm) * EMU_PER_CM)),
    )
    el = etree.fromstring(xml, _XP)
    slide.shapes._spTree.append(el)
    return sid


def add_replay_animation(slide, shape_ids: list[int], duration_ms: int, gap_ms: int = 0,
                         ease_in: float = 0.0, ease_out: float = 0.0, ease_mode: str = "keyframes"):
    """shape_ids 순서대로: 첫 번째는 클릭 시, 나머지는 '이전 효과 다음에'.
    ease_mode: keyframes = drawProgress 키프레임으로 이징 곡선을 직접 기록 (확실)
               attr      = cTn accel/decel 속성만 사용 (PowerPoint 경로 이동과 같은 방식)"""
    if ease_mode == "attr":
        accel_attrs = ""
        if ease_in > 0:
            accel_attrs += f' accel="{int(round(ease_in * 100000))}"'
        if ease_out > 0:
            accel_attrs += f' decel="{int(round(ease_out * 100000))}"'
        tavs = _tav_xml([(0.0, 0.0), (1.0, 1.0)])
    else:
        accel_attrs = ""
        curve = ease_curve(ease_in, ease_out) if (ease_in > 0 or ease_out > 0) else [(0.0, 0.0), (1.0, 1.0)]
        tavs = _tav_xml(curve)
    sld = slide._element
    timing = sld.find("{%s}timing" % NS["p"])
    if timing is None:
        timing = etree.fromstring(TIMING_XML.format(**NS), _XP)
        # p:sld 자식 순서: cSld, clrMapOvr, transition, timing, extLst
        ext = sld.find("{%s}extLst" % NS["p"])
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
    # 클릭 그룹(cTn id=4)의 childTnLst 에 효과들을 넣는다
    holder = timing.xpath(".//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par/p:cTn/p:childTnLst/p:par/p:cTn/p:childTnLst",
                          namespaces=NS)[-1]
    next_id = max(int(v) for v in timing.xpath(".//p:cTn/@id", namespaces=NS)) + 1
    for i, sid in enumerate(shape_ids):
        eff = etree.fromstring(INK_EFFECT_XML.format(
            **NS, id0=next_id, id1=next_id + 1, id2=next_id + 2, sid=sid, dur=duration_ms,
            node_type="clickEffect" if i == 0 else "afterEffect",
            delay=0 if i == 0 else gap_ms,
            accel_attrs=accel_attrs, tavs=tavs,
        ), _XP)
        # afterEffect 는 이전 효과가 끝난 뒤에 시작해야 하므로 별도 par 로 감싸고 delay 를 누적
        if i > 0:
            eff.find("p:cTn", NS).find("p:stCondLst/p:cond", NS).set("delay", str(i * duration_ms + i * gap_ms))
        holder.append(eff)
        next_id += 3


# ----------------------------------------------------------------------------
# 5. CLI
# ----------------------------------------------------------------------------
@dataclass
class ConvertOptions:
    template: Path | None = None
    width_cm: float | None = None
    height_cm: float | None = None
    duration: float = 2.0
    stroke_width: float | None = None
    color: str | None = None
    step: float | None = None
    split: str = "none"
    keep_png: bool = False
    ease_in: float = 0.0
    ease_out: float = 0.0
    ease_mode: str = "keyframes"
    timing: str = "channel"
    path_order: list[int] | None = None      # path 재생 순서 (원래 order 번호 목록)
    reversed_paths: set[int] | None = None   # 방향 뒤집을 path (원래 order 번호)


@dataclass
class ConvertResult:
    output: Path
    ink_objects: int
    strokes: int
    width_cm: float
    height_cm: float


def convert(svg_file: Path, out: Path, opt: ConvertOptions) -> ConvertResult:
    """SVG -> pptx 변환 본체. CLI 와 GUI 가 공용으로 쓴다."""
    if opt.template:
        prs = Presentation(str(opt.template))
    else:
        prs = Presentation()
        prs.slide_width = Emu(12_192_000)  # 16:9, 33.867 x 19.05 cm
        prs.slide_height = Emu(6_858_000)
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    slide = prs.slides.add_slide(blank)

    slide_w_cm = prs.slide_width / EMU_PER_CM
    slide_h_cm = prs.slide_height / EMU_PER_CM
    width_cm = opt.width_cm or slide_w_cm * 0.6
    height_cm = opt.height_cm or slide_h_cm * 0.8

    strokes = load_strokes(svg_file, opt.step, None)
    if opt.path_order or opt.reversed_paths:
        strokes = apply_order(strokes, opt.path_order, opt.reversed_paths)

    # 전체 bbox 기준으로 스케일/위치를 먼저 잡고, split 시에도 같은 변환을 쓴다
    whole = build_inkml(strokes, width_cm, height_cm, opt.stroke_width, opt.color, timing=opt.timing)
    off_x = int(round((slide_w_cm - whole.width_cm) / 2 * EMU_PER_CM))
    off_y = int(round((slide_h_cm - whole.height_cm) / 2 * EMU_PER_CM))

    png_files: list[Path] = []
    shape_ids: list[int] = []
    if opt.split == "none":
        png = out.with_suffix(".preview.png")
        margin = render_fallback_png(whole, png)
        png_files.append(png)
        shape_ids.append(add_ink_to_slide(slide, whole, png, margin, off_x, off_y, name="Ink 1"))
    else:
        # path 별로 나누되, 좌표계는 전체 bbox 기준 유지
        xs = [x for s in strokes for x, _ in s.points]
        ys = [y for s in strokes for _, y in s.points]
        minx, miny = min(xs), min(ys)
        scale = whole.width_cm / max(max(xs) - minx, 1e-9)  # cm per SVG unit (whole 과 동일)
        groups: dict[int, list[Stroke]] = {}
        for s in strokes:
            groups.setdefault(s.order, []).append(s)
        for k, (order, grp) in enumerate(sorted(groups.items())):
            gxs = [x for s in grp for x, _ in s.points]
            gys = [y for s in grp for _, y in s.points]
            part = build_inkml(grp, 0, None, opt.stroke_width, opt.color, timing=opt.timing, scale=scale)
            gx_off = off_x + int(round((min(gxs) - minx) * scale * EMU_PER_CM))
            gy_off = off_y + int(round((min(gys) - miny) * scale * EMU_PER_CM))
            png = out.with_suffix(f".preview{k+1}.png")
            margin = render_fallback_png(part, png)
            png_files.append(png)
            shape_ids.append(add_ink_to_slide(slide, part, png, margin, gx_off, gy_off, name=f"Ink {k+1}"))

    per = int(round(opt.duration * 1000 / max(1, len(shape_ids))))
    add_replay_animation(slide, shape_ids, per, ease_in=opt.ease_in, ease_out=opt.ease_out,
                         ease_mode=opt.ease_mode)

    prs.save(str(out))
    if not opt.keep_png:
        for p in png_files:
            p.unlink(missing_ok=True)
    return ConvertResult(output=out, ink_objects=len(shape_ids), strokes=len(strokes),
                         width_cm=whole.width_cm, height_cm=whole.height_cm)


def main(argv=None):
    ap = argparse.ArgumentParser(description="SVG -> PowerPoint 잉크 재생 애니메이션")
    ap.add_argument("svg", type=Path)
    ap.add_argument("-o", "--output", type=Path, help="출력 pptx (기본: svg 이름.pptx)")
    ap.add_argument("--template", type=Path, help="기존 pptx 를 바탕으로 마지막에 슬라이드 추가")
    ap.add_argument("--width-cm", type=float, default=None, help="로고 가로 크기 cm (기본: 슬라이드 폭의 60%%)")
    ap.add_argument("--height-cm", type=float, default=None, help="최대 세로 크기 cm (넘으면 축소)")
    ap.add_argument("--duration", type=float, default=2.0, help="그려지는 시간(초), 기본 2")
    ap.add_argument("--stroke-width", type=float, default=None, help="펜 굵기 cm (기본: SVG stroke-width 비례)")
    ap.add_argument("--color", default=None, help="모든 획 색 강제 (#RRGGBB)")
    ap.add_argument("--step", type=float, default=None, help="샘플 간격 (SVG 단위, 작을수록 정밀)")
    ap.add_argument("--split", choices=["none", "path"], default="none",
                    help="path: SVG의 path 마다 잉크 개체를 나눠 순서대로 재생")
    ap.add_argument("--keep-png", action="store_true", help="미리보기 PNG 를 남긴다")
    ap.add_argument("--ease-in", type=float, default=0.0,
                    help="부드럽게 시작: 전체 시간 중 가속 구간 비율 0~1 (예 0.3)")
    ap.add_argument("--ease-out", type=float, default=0.0,
                    help="부드럽게 끝: 전체 시간 중 감속 구간 비율 0~1 (예 0.3)")
    ap.add_argument("--ease-mode", choices=["keyframes", "attr"], default="keyframes",
                    help="keyframes: 키프레임으로 직접 기록(기본) / attr: accel,decel 속성만")
    ap.add_argument("--timing", choices=["none", "offset", "channel"], default="channel",
                    help="획 속도 정보 방식 (기본 channel: 길이에 비례해 일정 속도)")
    args = ap.parse_args(argv)

    out = args.output or args.svg.with_suffix(".pptx")
    opt = ConvertOptions(template=args.template, width_cm=args.width_cm, height_cm=args.height_cm,
                         duration=args.duration, stroke_width=args.stroke_width, color=args.color,
                         step=args.step, split=args.split, keep_png=args.keep_png,
                         ease_in=args.ease_in, ease_out=args.ease_out, ease_mode=args.ease_mode,
                         timing=args.timing)
    r = convert(args.svg, out, opt)
    print(f"완료: {r.output}  (잉크 개체 {r.ink_objects}개, 획 {r.strokes}개, "
          f"{r.width_cm:.2f} x {r.height_cm:.2f} cm, {args.duration}s)")


if __name__ == "__main__":
    main()
